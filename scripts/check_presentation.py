"""Layout check for the presentation deck.

    python scripts/check_presentation.py docs/AI_Statistician_Presentation.pptx

A rendered screenshot is the usual way to catch layout defects, but rendering a
.pptx needs LibreOffice, which is not always installed.  This measures the
shapes directly instead, for the three defects a render is looked at for:

  overflow      text wrapped against real glyph widths at its declared size,
                compared with the height of its box
  overlap       text boxes intruding on one another
  margins       content outside the slide, or inside the 0.5in safe edge --
                except a slide number, which is margin furniture by definition

It also checks that every embedded picture keeps its source aspect ratio, since
a stretched chart is a quieter defect than an overflowing one and just as wrong.

Text is measured in Arial, which is wider than Calibri at a given size, so the
estimate errs toward reporting overflow rather than hiding it.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

SLIDE_W, SLIDE_H, SAFE = 13.3, 7.5, 0.5
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
_fonts: dict[int, ImageFont.FreeTypeFont] = {}


def font(px: int):
    if px not in _fonts:
        _fonts[px] = ImageFont.truetype(FONT, px)
    return _fonts[px]


def inches(v) -> float:
    return Emu(v).inches if v is not None else 0.0


def needed_height(text: str, pt: float, width_in: float,
                  lead_pt: float | None) -> float:
    if not text.strip():
        return 0.0
    f = font(max(6, round(pt * 96 / 72)))
    usable = max(10.0, width_in * 96 - 10)
    lines = 0
    for para in text.split("\n"):
        words, cur, extra = para.split(), "", 0
        if not words:
            lines += 1
            continue
        for w in words:
            trial = f"{cur} {w}".strip()
            if f.getlength(trial) <= usable or not cur:
                cur = trial
            else:
                extra += 1
                cur = w
        lines += extra + 1
    return lines * (lead_pt or pt * 1.22) / 72.0


def text_of(sh):
    if not sh.has_text_frame:
        return "", 0.0, None
    parts, size, lead = [], 0.0, None
    for para in sh.text_frame.paragraphs:
        parts.append("".join(r.text for r in para.runs))
        ls = para.line_spacing
        if ls is not None and not isinstance(ls, float):
            lead = Emu(ls).pt
        for r in para.runs:
            if r.font.size:
                size = max(size, r.font.size.pt)
    return "\n".join(parts), size or 12.0, lead


def overlap(a, b) -> float:
    ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return ox * oy if ox > 0.02 and oy > 0.02 else 0.0


def main() -> int:
    prs = Presentation(sys.argv[1])
    issues = 0
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            x, y, w, h = (inches(sh.left), inches(sh.top),
                          inches(sh.width), inches(sh.height))
            text, pt, lead = text_of(sh)

            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 \
                    or y + h > SLIDE_H + 0.01:
                print(f"  slide {i:2d}  OFF-SLIDE   {text[:40]!r}")
                issues += 1
            elif text.strip() and not (text.strip().isdigit()
                                       and len(text.strip()) <= 3
                                       and y > SLIDE_H - SAFE - 0.2) \
                    and (x < SAFE - 0.01 or y < SAFE - 0.01
                                   or x + w > SLIDE_W - SAFE + 0.01
                                   or y + h > SLIDE_H - SAFE + 0.01):
                print(f"  slide {i:2d}  MARGIN      {text[:40]!r}")
                issues += 1

            if text.strip():
                need = needed_height(text, pt, w, lead)
                if need > h + 0.06:
                    print(f"  slide {i:2d}  OVERFLOW    needs {need:.2f}in, "
                          f"has {h:.2f}in  {text[:40]!r}")
                    issues += 1
                texts.append(((x, y, w, h), text))

            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                iw, ih = Image.open(io.BytesIO(sh.image.blob)).size
                src, drawn = iw / ih, w / h
                if abs(src - drawn) / src > 0.02:
                    print(f"  slide {i:2d}  DISTORTED   picture drawn at "
                          f"{drawn:.3f}, source is {src:.3f}")
                    issues += 1

        for a in range(len(texts)):
            for b in range(a + 1, len(texts)):
                ra, ta = texts[a]
                rb, tb = texts[b]
                if (ra[2] < 0.7 and ra[3] < 0.7) or (rb[2] < 0.7 and rb[3] < 0.7):
                    continue            # a letter centred on its own disc
                if overlap(ra, rb) > 0.05:
                    print(f"  slide {i:2d}  OVERLAP     {ta[:22]!r} / {tb[:22]!r}")
                    issues += 1

    print(f"\n  {len(prs.slides)} slides, {issues} issue(s)")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
